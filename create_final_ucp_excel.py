import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

def generate_final_ucp_excel():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Consolidado UCP & Estimativa"
    ws.views.sheetView[0].showGridLines = True

    # Color Palette
    COLOR_HEADER_BG = "0F172A"       # Dark Slate
    COLOR_HEADER_FG = "FFFFFF"       # White
    COLOR_ACCENT = "0D9488"          # Teal
    COLOR_ZEBRA = "F8FAFC"           # Light slate
    COLOR_SUBTOTAL_BG = "F0FDF4"     # Light green
    COLOR_SUBTOTAL_FG = "166534"     # Dark green
    COLOR_BORDER = "CBD5E1"          # Gray border
    COLOR_CARD_BG = "F1F5F9"        # Slate light

    # Fonts
    font_title = Font(name="Arial", size=16, bold=True, color="0F172A")
    font_subtitle = Font(name="Arial", size=11, bold=True, color="0D9488")
    font_sec = Font(name="Arial", size=12, bold=True, color="0F172A")
    font_th = Font(name="Arial", size=10, bold=True, color=COLOR_HEADER_FG)
    font_td = Font(name="Arial", size=10, color="1E293B")
    font_td_bold = Font(name="Arial", size=10, bold=True, color="1E293B")
    font_highlight = Font(name="Arial", size=11, bold=True, color=COLOR_SUBTOTAL_FG)
    font_badge = Font(name="Arial", size=13, bold=True, color="FFFFFF")
    font_total_label = Font(name="Arial", size=11, bold=True, color="0F172A")
    font_total_val = Font(name="Arial", size=11, bold=True, color=COLOR_SUBTOTAL_FG)

    # Fills
    fill_th = PatternFill(start_color=COLOR_HEADER_BG, end_color=COLOR_HEADER_BG, fill_type="solid")
    fill_zebra = PatternFill(start_color=COLOR_ZEBRA, end_color=COLOR_ZEBRA, fill_type="solid")
    fill_subtotal = PatternFill(start_color=COLOR_SUBTOTAL_BG, end_color=COLOR_SUBTOTAL_BG, fill_type="solid")
    fill_accent = PatternFill(start_color=COLOR_ACCENT, end_color=COLOR_ACCENT, fill_type="solid")
    fill_card = PatternFill(start_color=COLOR_CARD_BG, end_color=COLOR_CARD_BG, fill_type="solid")

    thin_border = Border(
        left=Side(style='thin', color=COLOR_BORDER),
        right=Side(style='thin', color=COLOR_BORDER),
        top=Side(style='thin', color=COLOR_BORDER),
        bottom=Side(style='thin', color=COLOR_BORDER)
    )

    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center", wrap_text=True)
    align_right = Alignment(horizontal="right", vertical="center")

    # Title
    ws['A1'] = "CONSOLIDADO UCP: AUCP, ESFORÇO, CRONOGRAMA E CUSTOS"
    ws['A1'].font = font_title
    ws['A2'] = "Método de Pontos por Casos de Uso (Karner / Schneider & Winters) — HRTech Core"
    ws['A2'].font = font_subtitle

    ws.row_dimensions[1].height = 25
    ws.row_dimensions[2].height = 20

    # 1. Inputs Block
    ws['A4'] = "1. DADOS DE ENTRADA DO PROJETO"
    ws['A4'].font = font_sec

    inputs = [
        ("Sistema:", "HRTech Core — Sistema Modular de Gestão de RH (LPS)"),
        ("Pontos de Casos de Uso Não Ajustados (UUCP):", 178),
        ("Fator de Complexidade Técnica (TCF):", 1.105),
        ("Fator de Complexidade Ambiental (ECF):", 0.755),
        ("Valor da Homem-Hora (R$):", 50.00),
        ("Carga Horária Semanal por Dev (horas):", 40)
    ]

    for idx, (label, val) in enumerate(inputs, start=5):
        ws[f'A{idx}'] = label
        ws[f'A{idx}'].font = Font(name="Arial", size=9, bold=True, color="475569")
        c_val = ws[f'B{idx}']
        c_val.value = val
        c_val.font = font_td_bold
        ws.row_dimensions[idx].height = 19
        if "R$" in label:
            c_val.number_format = 'R$ #,##0.00'
        elif isinstance(val, float):
            c_val.number_format = '0.000'

    # 2. AUCP Calculation
    ws['A12'] = "2. CÁLCULO DO UCP AJUSTADO (AUCP)"
    ws['A12'].font = font_sec

    ws['A13'] = "Fórmula:"
    ws['B13'] = "AUCP = UUCP × TCF × ECF"
    ws['A14'] = "Substituição:"
    ws['B14'] = '=CONCATENATE("AUCP = ", B6, " × ", TEXT(B7,"0.000"), " × ", TEXT(B8,"0.000"))'
    ws['A15'] = "AUCP RESULTANTE:"
    ws['B15'] = "=B6*B7*B8"

    ws['A13'].font = Font(name="Arial", size=9, bold=True, color="475569")
    ws['B13'].font = Font(name="Arial", size=10, italic=True)
    ws['A14'].font = Font(name="Arial", size=9, bold=True, color="475569")
    ws['B14'].font = font_td
    ws['A15'].font = Font(name="Arial", size=11, bold=True, color="0F172A")
    
    ws['B15'].font = font_badge
    ws['B15'].fill = fill_accent
    ws['B15'].alignment = align_center
    ws['B15'].number_format = '0.000'
    ws.row_dimensions[15].height = 28

    # 3. Environmental Rules (Schneider & Winters)
    ws['A17'] = "3. AVALIAÇÃO DA TAXA DE PRODUTIVIDADE (REGRA DE SCHNEIDER & WINTERS)"
    ws['A17'].font = font_sec

    env_headers = ["Variável", "Regra / Condição", "Quantidade Apurada", "Taxa Recomendada"]
    ws.row_dimensions[18].height = 24
    for col_idx, h in enumerate(env_headers, start=1):
        cell = ws.cell(row=18, column=col_idx, value=h)
        cell.font = font_th
        cell.fill = fill_th
        cell.alignment = align_center
        cell.border = thin_border

    ws.cell(row=19, column=1, value="Fatores F1-F6 < 3 (X)").font = font_td_bold
    ws.cell(row=19, column=2, value="Qtd de fatores de processo/equipe com nota < 3").font = font_td
    ws.cell(row=19, column=3, value=0).font = font_td_bold
    ws.cell(row=19, column=3).alignment = align_center
    ws.cell(row=19, column=4, value="—").alignment = align_center

    ws.cell(row=20, column=1, value="Fatores F7-F8 > 3 (Y)").font = font_td_bold
    ws.cell(row=20, column=2, value="Qtd de fatores de risco (parcial/dificuldade) com nota > 3").font = font_td
    ws.cell(row=20, column=3, value=0).font = font_td_bold
    ws.cell(row=20, column=3).alignment = align_center
    ws.cell(row=20, column=4, value="—").alignment = align_center

    ws.cell(row=21, column=1, value="Total (X + Y)").font = font_total_label
    ws.cell(row=21, column=2, value="Condição: (X + Y) <= 2 → 20 h/UCP").font = font_td_bold
    ws.cell(row=21, column=3, value="=C19+C20").font = font_total_val
    ws.cell(row=21, column=3).alignment = align_center
    ws.cell(row=21, column=4, value="20 h / UCP").font = font_total_val
    ws.cell(row=21, column=4).alignment = align_center
    ws.cell(row=21, column=4).fill = fill_subtotal

    for r in range(19, 22):
        ws.row_dimensions[r].height = 22
        for c in range(1, 5):
            ws.cell(row=r, column=c).border = thin_border

    # 4. Comparative Table of Productivity & Schedule
    ws['A23'] = "4. TABELA COMPARATIVA DE PRODUTIVIDADE E CRONOGRAMA"
    ws['A23'].font = font_sec

    comp_headers = ["Produtividade (h/UCP)", "Esforço Total (Horas)", "Custo Total (R$)", "1 Desenvolvedor", "2 Desenvolvedores", "4 Desenvolvedores"]
    ws.row_dimensions[24].height = 26
    for col_idx, h in enumerate(comp_headers, start=1):
        cell = ws.cell(row=24, column=col_idx, value=h)
        cell.font = font_th
        cell.fill = fill_th
        cell.alignment = align_center
        cell.border = thin_border

    scenarios = [
        (15, "15 h/UCP (Otimista)"),
        (20, "20 h/UCP (OFICIAL RECOMENDADO)"),
        (25, "25 h/UCP (Conservador)"),
        (28, "28 h/UCP (Contingência / Alto Risco)")
    ]

    for idx, (rate, label) in enumerate(scenarios, start=25):
        ws.row_dimensions[idx].height = 24
        
        c_rate = ws.cell(row=idx, column=1, value=label)
        c_hours = ws.cell(row=idx, column=2, value=f"=$B$15*{rate}")
        c_cost = ws.cell(row=idx, column=3, value=f"=B{idx}*$B$9")
        
        c_dev1 = ws.cell(row=idx, column=4, value=f"=B{idx}/(1*$B$10)")
        c_dev2 = ws.cell(row=idx, column=5, value=f"=B{idx}/(2*$B$10)")
        c_dev4 = ws.cell(row=idx, column=6, value=f"=B{idx}/(4*$B$10)")

        c_rate.font = font_td_bold if rate == 20 else font_td
        c_hours.font = font_highlight if rate == 20 else font_td
        c_cost.font = font_highlight if rate == 20 else font_td
        
        c_hours.number_format = '#,##0.0 "h"'
        c_cost.number_format = 'R$ #,##0.00'

        for c_dev in [c_dev1, c_dev2, c_dev4]:
            c_dev.font = font_td
            c_dev.number_format = '0.0 "sem"'
            c_dev.alignment = align_center

        c_rate.alignment = align_left
        c_hours.alignment = align_center
        c_cost.alignment = align_center

        if rate == 20:
            for col in range(1, 7):
                ws.cell(row=idx, column=col).fill = fill_subtotal

        for col in range(1, 7):
            ws.cell(row=idx, column=col).border = thin_border

    # 5. Financial Summary Box
    ws['A30'] = "5. ORÇAMENTO FINANCEIRO E RESUMO DA ESTIMATIVA"
    ws['A30'].font = font_sec

    summary_rows = [
        ("UCP Ajustado (AUCP):", "=B15", "0.000 AUCP"),
        ("Taxa de Esforço Oficial:", 20, "20 homem-horas / UCP"),
        ("Esforço Total Estimado:", "=B26", "#,##0.0 horas"),
        ("Valor da Homem-Hora:", "=B9", "R$ #,##0.00"),
        ("INVESTIMENTO TOTAL ESTIMADO:", "=C26", "R$ #,##0.00")
    ]

    for idx, (lbl, val, fmt) in enumerate(summary_rows, start=31):
        ws.row_dimensions[idx].height = 22
        c_l = ws.cell(row=idx, column=1, value=lbl)
        c_v = ws.cell(row=idx, column=2, value=val)

        c_l.font = Font(name="Arial", size=10, bold=True, color="0F172A")
        c_v.font = font_highlight if "TOTAL" in lbl else font_td_bold

        if "TOTAL" in lbl:
            c_v.font = Font(name="Arial", size=12, bold=True, color=COLOR_SUBTOTAL_FG)
            c_v.fill = fill_subtotal

        if "R$" in fmt:
            c_v.number_format = 'R$ #,##0.00'
        elif "horas" in fmt:
            c_v.number_format = '#,##0.0 "horas"'
        elif "AUCP" in fmt:
            c_v.number_format = '0.000'

        c_l.alignment = align_left
        c_v.alignment = align_left

    # Column Widths
    col_widths = {
        'A': 38,
        'B': 45,
        'C': 26,
        'D': 20,
        'E': 20,
        'F': 20
    }
    for col_l, w in col_widths.items():
        ws.column_dimensions[col_l].width = w

    output_path = "/home/fernando/Documentos/Faculdade/Projeto de medição e analise/Calculo_Final_UCP_AUCP_HRTech_Core.xlsx"
    wb.save(output_path)
    print(f"Final UCP Excel generated at: {output_path}")

generate_final_ucp_excel()
