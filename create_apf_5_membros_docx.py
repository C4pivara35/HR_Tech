import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def set_table_borders(table, color="D1D5DB", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideV w:val="none"/>'
        f'<w:left w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def format_cell(cell, text, bold=False, italic=False, color=RGBColor(31, 41, 55), font_size=10, align=WD_ALIGN_PARAGRAPH.LEFT, bg_color=None):
    if bg_color:
        set_cell_background(cell, bg_color)
    set_cell_margins(cell)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(str(text))
    run.font.name = 'Arial'
    run.font.size = Pt(font_size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color

def add_heading_1(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(16)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = 'Arial'
    run.font.size = Pt(14)
    run.font.bold = True
    run.font.color.rgb = RGBColor(30, 58, 138)

def add_heading_2(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = 'Arial'
    run.font.size = Pt(11.5)
    run.font.bold = True
    run.font.color.rgb = RGBColor(17, 24, 39)

def add_body_p(doc, text, bold_prefix="", italic=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_bold = p.add_run(bold_prefix)
        r_bold.font.name = 'Arial'
        r_bold.font.size = Pt(10.5)
        r_bold.font.bold = True
        r_bold.font.color.rgb = RGBColor(31, 41, 55)
    run = p.add_run(text)
    run.font.name = 'Arial'
    run.font.size = Pt(10.5)
    run.font.italic = italic
    run.font.color.rgb = RGBColor(55, 65, 81)

def build_student_docx(filename):
    doc = Document()

    # Page Margins (Normal 2.5cm)
    for section in doc.sections:
        section.top_margin = Inches(0.98)
        section.bottom_margin = Inches(0.98)
        section.left_margin = Inches(0.98)
        section.right_margin = Inches(0.98)

    # Header / Top Title (Student style - direct & clean)
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(4)
    run_t = p_title.add_run("Trabalho de Medição e Análise de Software")
    run_t.font.name = 'Arial'
    run_t.font.size = Pt(18)
    run_t.font.bold = True
    run_t.font.color.rgb = RGBColor(30, 58, 138)

    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_after = Pt(12)
    run_s = p_sub.add_run("Atividade em Grupo: Contagem de Pontos de Função (IFPUG) — ALIs e AIEs")
    run_s.font.name = 'Arial'
    run_s.font.size = Pt(12)
    run_s.font.bold = True
    run_s.font.color.rgb = RGBColor(75, 85, 99)

    # Info Box (Simple student header box)
    tbl_info = doc.add_table(rows=4, cols=2)
    tbl_info.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl_info, color="E5E7EB", sz="4")

    info_data = [
        ("Disciplina:", "Medição e Análise de Software — PUCPR"),
        ("Sistema:", "HRTech Core — Plataforma de Gestão de RH"),
        ("Equipe (5 membros):", "Fernando Lopes Duarte + 4 integrantes da equipe"),
        ("Regra da Atividade:", "2 ALIs + 2 AIEs por integrante = 10 ALIs + 10 AIEs (Total: 20 arquivos)")
    ]

    for row_idx, (label, val) in enumerate(info_data):
        row = tbl_info.rows[row_idx]
        format_cell(row.cells[0], label, bold=True, color=RGBColor(30, 58, 138), bg_color="F3F4F6")
        format_cell(row.cells[1], val, bold=False, color=RGBColor(31, 41, 55), bg_color="F9FAFB")
        row.cells[0].width = Inches(2.0)
        row.cells[1].width = Inches(4.5)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    add_body_p(doc, "Professor, a nossa equipe é formada por 5 integrantes. Conforme a regra orientada em aula (2 ALIs e 2 AIEs por aluno), organizamos a contagem dos 20 arquivos de dados do nosso sistema HRTech Core divididos entre os membros do grupo. Abaixo mostramos a divisão das tarefas, o detalhamento de cada arquivo (RLR, DER e complexidade) e a tabela final com o total de Pontos de Função.")

    # Section 1: Divisão da Equipe
    add_heading_1(doc, "1. Divisão do Trabalho no Grupo (5 Integrantes)")

    tbl_group = doc.add_table(rows=6, cols=4)
    tbl_group.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl_group, color="D1D5DB")

    headers_g = ["Integrante", "Módulo do Sistema", "2 ALIs (Internos)", "2 AIEs (Externos)"]
    hdr_row = tbl_group.rows[0]
    for i, h in enumerate(headers_g):
        format_cell(hdr_row.cells[i], h, bold=True, color=RGBColor(255, 255, 255), bg_color="1E3A8A", align=WD_ALIGN_PARAGRAPH.CENTER)

    group_rows = [
        ("Membro 1", "Cadastro & Organograma", "ALI 1: ALI_COLABORADORES\nALI 2: ALI_ESTRUTURA_ORGANIZACIONAL", "AIE 1: AIE_RECEITA_FEDERAL_CPF\nAIE 2: AIE_GOV_ESOCIAL"),
        ("Membro 2", "Ponto & Banco de Horas", "ALI 3: ALI_REGISTROS_PONTO\nALI 4: ALI_BANCO_HORAS_ACUMULADO", "AIE 3: AIE_RELOGIO_PONTO_REP\nAIE 4: AIE_API_NTP_HORA_CERTA"),
        ("Membro 3", "Férias & Benefícios", "ALI 5: ALI_SOLICITACOES_FERIAS\nALI 6: ALI_CARTOES_BENEFICIOS", "AIE 5: AIE_OPERADORA_VR_VA\nAIE 6: AIE_OPERADORA_SAUDE"),
        ("Membro 4", "EPIs & Exames ASO", "ALI 7: ALI_EPIS_EQUIPAMENTOS\nALI 8: ALI_EXAMES_ASO", "AIE 7: AIE_MINISTERIO_TRABALHO_CA\nAIE 8: AIE_CLINICA_MEDICA_OCUPACIONAL"),
        ("Membro 5", "Auditoria & Portal Corretor", "ALI 9: ALI_AUDITORIA_LOGS\nALI 10: ALI_APOLICES_CORRETOR", "AIE 9: AIE_SEGURADORA_PORTAL\nAIE 10: AIE_SISTEMA_BANCARIO_PIX"),
    ]

    for r_idx, data in enumerate(group_rows, start=1):
        row = tbl_group.rows[r_idx]
        bg = "FFFFFF" if r_idx % 2 != 0 else "F9FAFB"
        format_cell(row.cells[0], data[0], bold=True, color=RGBColor(30, 58, 138), bg_color=bg)
        format_cell(row.cells[1], data[1], bold=False, bg_color=bg)
        format_cell(row.cells[2], data[2], bold=False, bg_color=bg)
        format_cell(row.cells[3], data[3], bold=False, bg_color=bg)

    col_widths_g = [Inches(1.2), Inches(1.8), Inches(2.2), Inches(2.2)]
    for row in tbl_group.rows:
        for idx, w in enumerate(col_widths_g):
            row.cells[idx].width = w

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # Section 2: Detalhamento por Membro
    add_heading_1(doc, "2. Detalhamento dos Arquivos (ALIs e AIEs)")

    members_details = [
        ("Membro 1 — Cadastro & Organograma", [
            ("ALI 1: ALI_COLABORADORES", "ALI", "Média (10 PF)", "Tabela principal dos colaboradores da empresa com dados pessoais e contratuais.", "4 (Colaborador, Endereço, Contato, DadosBancarios)", "24 (id, matricula, nome, cpf, rg, data_nasc, salario, email, telefone, cep, logradouro, numero, bairro, cidade, uf, banco, agencia, conta, pix, etc.)"),
            ("ALI 2: ALI_ESTRUTURA_ORGANIZACIONAL", "ALI", "Baixa (7 PF)", "Estrutura dos departamentos, cargos e níveis da empresa.", "3 (Departamento, Cargo, NivelHierarquico)", "14 (id, nome_depto, centro_custo, gestor_id, titulo_cargo, nivel, cbo, descricao, faixa_sal_inicio, faixa_sal_fim, etc.)"),
            ("AIE 1: AIE_RECEITA_FEDERAL_CPF", "AIE", "Baixa (5 PF)", "Consulta externa via API na Receita Federal para validar CPF/CNPJ.", "2 (CadastroReceita, SituacaoCadastral)", "12 (cpf_cnpj, nome, situacao, data_nasc, data_consulta, protocolo, digito_verificador, uf, hash, etc.)"),
            ("AIE 2: AIE_GOV_ESOCIAL", "AIE", "Média (7 PF)", "Integração externa com governo federal para envio do evento eSocial S-2200 (Admissão).", "2 (EventoAdmissao, TrabalhadorEsocial)", "35 (cpf, nis, matricula_esocial, data_admissao, cbo, remuneracao, carga_horaria, recibo_entrega, etc.)")
        ]),
        ("Membro 2 — Ponto & Banco de Horas", [
            ("ALI 3: ALI_REGISTROS_PONTO", "ALI", "Baixa (7 PF)", "Marcações de ponto eletrônico com geolocalização e carimbo de hora.", "3 (RegistroPonto, Geolocalizacao, Dispositivo)", "18 (id, colaborador_id, data_hora, tipo, ip, latitude, longitude, precisao, dispositivo, hash_sha256, etc.)"),
            ("ALI 4: ALI_BANCO_HORAS_ACUMULADO", "ALI", "Baixa (7 PF)", "Saldos acumulados de horas extras e compensações de ponto.", "3 (SaldoBancoHoras, HorasExtras, ExtratoCompensacao)", "15 (id, colaborador_id, mes, ano, saldo_anterior, horas_credito, horas_debito, saldo_atual, status, etc.)"),
            ("AIE 3: AIE_RELOGIO_PONTO_REP", "AIE", "Baixa (5 PF)", "Importação de arquivo AFD dos relógios de ponto físicos da empresa.", "2 (BilheteAFD, CabecalhoAFD)", "10 (numero_rep, cnpj_empresa, nsr, data_marcacao, hora_marcacao, pis, checksum, status, etc.)"),
            ("AIE 4: AIE_API_NTP_HORA_CERTA", "AIE", "Baixa (5 PF)", "Servidor externo de hora oficial NTP.br para validar o horário do ponto.", "1 (RespostaNTP)", "6 (timestamp_ntp, stratum, precision_ms, delay, dispersion, server_ip)")
        ]),
        ("Membro 3 — Férias & Benefícios Corporativos", [
            ("ALI 5: ALI_SOLICITACOES_FERIAS", "ALI", "Baixa (7 PF)", "Solicitações de férias dos funcionários e cálculo de 1/3 constitucional.", "3 (PeriodoAquisitivo, SolicitacaoFerias, AbonoPecuniario)", "16 (id, colaborador_id, inicio_aquisitivo, fim_aquisitivo, data_inicio, dias_gozo, dias_abono, adiantamento13, status, etc.)"),
            ("ALI 6: ALI_CARTOES_BENEFICIOS", "ALI", "Baixa (7 PF)", "Gestão dos cartões de VR/VA, transporte e plano de saúde.", "3 (CartaoBeneficio, SolicitacaoCarga, DependentePlano)", "18 (id, colaborador_id, tipo_beneficio, operadora, numero_cartao, valor_empresa, valor_desconto, status, etc.)"),
            ("AIE 5: AIE_OPERADORA_VR_VA", "AIE", "Baixa (5 PF)", "API da operadora de benefícios (Ticket/Sodexo/Alelo) para pedir recarga de cartão.", "2 (PedidoCarga, ExtratoOperadora)", "14 (cod_cliente, contrato, cpf_favorecido, valor_credito, data_carga, cod_pedido, status, protocolo, etc.)"),
            ("AIE 6: AIE_OPERADORA_SAUDE", "AIE", "Média (7 PF)", "Interface externa com a operadora de plano de saúde para inclusão de dependentes e coparticipação.", "3 (FaturaOperadora, MovimentacaoVidas, SinistroCoparticipacao)", "22 (cnpj_operadora, numero_apolice, cpf_titular, cpf_dependente, tipo_movimentacao, valor_mensalidade, etc.)")
        ]),
        ("Membro 4 — Segurança do Trabalho (EPIs & Exames ASO)", [
            ("ALI 7: ALI_EPIS_EQUIPAMENTOS", "ALI", "Baixa (7 PF)", "Ficha de controle de entregas de EPIs aos funcionários.", "3 (FichaEPI, ItemEquipamento, TermoResponsabilidade)", "15 (id, colaborador_id, nome_epi, ca_numero, data_entrega, quantidade, data_validade_epi, assinatura, etc.)"),
            ("ALI 8: ALI_EXAMES_ASO", "ALI", "Baixa (7 PF)", "Prontuário de ASO (Atestado de Saúde Ocupacional) e exames periódicos.", "3 (ProntuarioASO, ExameRealizado, MedicoExaminador)", "16 (id, colaborador_id, tipo_aso, data_exame, resultado, nome_medico, crm, uf_crm, validade_aso, etc.)"),
            ("AIE 7: AIE_MINISTERIO_TRABALHO_CA", "AIE", "Baixa (5 PF)", "Consulta ao banco público do MTE para verificar a validade do Certificado de Aprovação (CA) do EPI.", "1 (ConsultaCA)", "8 (numero_ca, nome_equipamento, fabricante, data_validade_ca, situacao, laudo_tecnico, etc.)"),
            ("AIE 8: AIE_CLINICA_MEDICA_OCUPACIONAL", "AIE", "Baixa (5 PF)", "Integração via API com a clínica parceira para agendar exames admissionais e receber laudos.", "2 (AgendamentoExame, LaudoDigital)", "12 (cnpj_clinica, codigo_agendamento, cpf_funcionario, tipo_exame, data_agendada, status, laudo_pdf_url, etc.)")
        ]),
        ("Membro 5 — Compliance, Auditoria & Portal do Corretor", [
            ("ALI 9: ALI_AUDITORIA_LOGS", "ALI", "Baixa (7 PF)", "Logs de auditoria imutáveis com hash para rastrear quem alterou dados no sistema.", "3 (LogOperacao, TrilhaLGPD, EventoSeguranca)", "14 (id, tenant_id, usuario_id, acao, entidade, dados_anteriores, dados_novos, ip, data_hora, hash_sha256, etc.)"),
            ("ALI 10: ALI_APOLICES_CORRETOR", "ALI", "Média (10 PF)", "Gestão das apólices exclusivas de seguro do Portal do Corretor (FinCorp).", "4 (ApoliceSeguro, CorretorParceiro, CoberturaContratada, SinistroRegistrado)", "26 (id, numero_apolice, corretor_id, premio_total, comissao_porcentagem, inicio_vigencia, fim_vigencia, status, etc.)"),
            ("AIE 9: AIE_SEGURADORA_PORTAL", "AIE", "Média (7 PF)", "Comunicação via API com a seguradora para emitir apólices e validar propostas.", "3 (PropostaSeguradora, EndossoApolice, ComissaoCorretor)", "20 (codigo_proposta, cnpj_seguradora, dados_estipulante, valor_premio_net, imposto_iof, status_emissao, etc.)"),
            ("AIE 10: AIE_SISTEMA_BANCARIO_PIX", "AIE", "Baixa (5 PF)", "Integração com gateway bancário para pagamento automático via PIX.", "2 (TransacaoPix, ComprovanteBancario)", "16 (end_to_end_id, txid, chave_pix, valor, data_hora, status, codigo_autenticacao, qr_code, etc.)")
        ]),
    ]

    for member_title, files in members_details:
        add_heading_2(doc, member_title)
        tbl_m = doc.add_table(rows=len(files) + 1, cols=6)
        tbl_m.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_table_borders(tbl_m, color="E5E7EB")

        headers_m = ["Arquivo", "Tipo", "Complexidade", "Descrição", "RLRs", "DERs"]
        hdr = tbl_m.rows[0]
        for i, h in enumerate(headers_m):
            format_cell(hdr.cells[i], h, bold=True, color=RGBColor(255, 255, 255), bg_color="3B82F6", align=WD_ALIGN_PARAGRAPH.CENTER)

        for r_idx, f_data in enumerate(files, start=1):
            row = tbl_m.rows[r_idx]
            bg = "FFFFFF" if r_idx % 2 != 0 else "F9FAFB"
            format_cell(row.cells[0], f_data[0], bold=True, color=RGBColor(30, 58, 138), bg_color=bg)
            format_cell(row.cells[1], f_data[1], bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, bg_color=bg)
            format_cell(row.cells[2], f_data[2], bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, bg_color=bg)
            format_cell(row.cells[3], f_data[3], bg_color=bg)
            format_cell(row.cells[4], f_data[4], bg_color=bg)
            format_cell(row.cells[5], f_data[5], bg_color=bg)

        col_w = [Inches(1.8), Inches(0.6), Inches(1.1), Inches(1.7), Inches(1.1), Inches(1.2)]
        for row in tbl_m.rows:
            for idx, w in enumerate(col_w):
                row.cells[idx].width = w

        doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # Section 3: Tabela Consolidada dos Pontos de Função
    add_heading_1(doc, "3. Resumo da Contagem Final de Pontos de Função")

    add_body_p(doc, "Somando a pontuação de todos os 20 arquivos contados pelos 5 integrantes do grupo, chegamos ao total de 132 Pontos de Função Não Ajustados:")

    tbl_total = doc.add_table(rows=7, cols=4)
    tbl_total.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl_total, color="9CA3AF")

    headers_t = ["Integrante do Grupo", "Pontos ALIs (Internos)", "Pontos AIEs (Externos)", "Total do Integrante"]
    hdr_t = tbl_total.rows[0]
    for i, h in enumerate(headers_t):
        format_cell(hdr_t.cells[i], h, bold=True, color=RGBColor(255, 255, 255), bg_color="1E3A8A", align=WD_ALIGN_PARAGRAPH.CENTER)

    totals_rows = [
        ("Membro 1 (Cadastro & Organograma)", "10 PF + 7 PF = 17 PF", "5 PF + 7 PF = 12 PF", "29 PF"),
        ("Membro 2 (Ponto & Banco de Horas)", "7 PF + 7 PF = 14 PF", "5 PF + 5 PF = 10 PF", "24 PF"),
        ("Membro 3 (Férias & Benefícios)", "7 PF + 7 PF = 14 PF", "5 PF + 7 PF = 12 PF", "26 PF"),
        ("Membro 4 (EPIs & Exames ASO)", "7 PF + 7 PF = 14 PF", "5 PF + 5 PF = 10 PF", "24 PF"),
        ("Membro 5 (Auditoria & Corretor)", "7 PF + 10 PF = 17 PF", "7 PF + 5 PF = 12 PF", "29 PF"),
    ]

    for r_idx, t_data in enumerate(totals_rows, start=1):
        row = tbl_total.rows[r_idx]
        bg = "FFFFFF" if r_idx % 2 != 0 else "F9FAFB"
        format_cell(row.cells[0], t_data[0], bold=True, color=RGBColor(31, 41, 55), bg_color=bg)
        format_cell(row.cells[1], t_data[1], align=WD_ALIGN_PARAGRAPH.CENTER, bg_color=bg)
        format_cell(row.cells[2], t_data[2], align=WD_ALIGN_PARAGRAPH.CENTER, bg_color=bg)
        format_cell(row.cells[3], t_data[3], bold=True, color=RGBColor(30, 58, 138), align=WD_ALIGN_PARAGRAPH.CENTER, bg_color=bg)

    # Total Row
    row_last = tbl_total.rows[6]
    format_cell(row_last.cells[0], "TOTAL DA EQUIPE", bold=True, color=RGBColor(255, 255, 255), bg_color="1E3A8A")
    format_cell(row_last.cells[1], "76 PF (10 ALIs)", bold=True, color=RGBColor(255, 255, 255), align=WD_ALIGN_PARAGRAPH.CENTER, bg_color="1E3A8A")
    format_cell(row_last.cells[2], "56 PF (10 AIEs)", bold=True, color=RGBColor(255, 255, 255), align=WD_ALIGN_PARAGRAPH.CENTER, bg_color="1E3A8A")
    format_cell(row_last.cells[3], "132 PF", bold=True, color=RGBColor(255, 255, 255), align=WD_ALIGN_PARAGRAPH.CENTER, bg_color="1E3A8A")

    col_w_t = [Inches(2.5), Inches(1.6), Inches(1.6), Inches(1.5)]
    for row in tbl_total.rows:
        for idx, w in enumerate(col_w_t):
            row.cells[idx].width = w

    doc.save(filename)
    print(f"Documento Word salvo com sucesso em: {filename}")

if __name__ == "__main__":
    build_student_docx("/home/fernando/Documentos/Faculdade/Projeto de medição e analise/Estimativa_APF_5_Membros_ALIs_AIEs_HRTech.docx")
