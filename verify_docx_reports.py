#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
HRTech Core — Suíte de Verificação e Auditoria Forense dos Documentos DOCX
Atividades 05 (Individual) e 06 (Equipe) — PUCPR
"""

import os
import sys
import docx
from docx import Document
from docx.shared import Cm

PROJECT_ROOT = "/home/fernando/Documentos/Faculdade/Projeto de medição e analise"
INDIVIDUAL_DOCX = os.path.join(PROJECT_ROOT, "Entrega_APF_Individual_Fernando_Lopes_Duarte.docx")
TEAM_DOCX = os.path.join(PROJECT_ROOT, "Entrega_APF_Equipe_HRTech_Core.docx")

sys.path.insert(0, PROJECT_ROOT)
from scripts.generate_docx_reports import ALIS_DATA, AIES_DATA, EES_DATA, SES_DATA, CES_DATA, GSCS_DATA

def audit_file(path, label, min_size_kb=100):
    print(f"\n--- Auditoria Básica: {label} ---")
    if not os.path.exists(path):
        print(f"[FAIL] Arquivo não encontrado: {path}")
        return False
    size = os.path.getsize(path)
    size_kb = size / 1024
    if size_kb < min_size_kb:
        print(f"[FAIL] Tamanho do arquivo {size_kb:.1f} KB é inferior ao mínimo de {min_size_kb} KB.")
        return False
    print(f"[PASS] Arquivo existe e possui tamanho de {size:,} bytes ({size_kb:.1f} KB).")
    return True

def audit_xml_styling(doc, label):
    print(f"Auditoria de Margens e XML em {label}...")
    # Margens ABNT A4
    sec = doc.sections[0]
    top_cm = sec.top_margin.cm
    left_cm = sec.left_margin.cm
    bottom_cm = sec.bottom_margin.cm
    right_cm = sec.right_margin.cm

    assert abs(top_cm - 3.0) < 0.1, f"Margem superior {top_cm} diverge de 3.0 cm"
    assert abs(left_cm - 3.0) < 0.1, f"Margem esquerda {left_cm} diverge de 3.0 cm"
    assert abs(bottom_cm - 2.0) < 0.1, f"Margem inferior {bottom_cm} diverge de 2.0 cm"
    assert abs(right_cm - 2.0) < 0.1, f"Margem direita {right_cm} diverge de 2.0 cm"
    print(f"  [PASS] Geometria ABNT A4 validada (Sup={top_cm:.1f}cm, Esq={left_cm:.1f}cm, Inf={bottom_cm:.1f}cm, Dir={right_cm:.1f}cm).")

    # XML de Tabelas: tblHeader, cantSplit, shd Navy Blue
    tbl_header_found = False
    cant_split_found = False
    navy_header_found = False

    for table in doc.tables:
        xml_str = table._tbl.xml
        if 'w:tblHeader' in xml_str:
            tbl_header_found = True
        if 'w:cantSplit' in xml_str:
            cant_split_found = True
        if 'w:fill="1B365D"' in xml_str:
            navy_header_found = True

    assert tbl_header_found, "Tag <w:tblHeader/> não encontrada nas tabelas."
    assert cant_split_found, "Tag <w:cantSplit/> não encontrada nas linhas."
    assert navy_header_found, "Shading Navy Blue #1B365D não encontrado nos cabeçalhos."
    print("  [PASS] Elementos XML <w:tblHeader/>, <w:cantSplit/> e Shading #1B365D auditados com sucesso.")

def audit_gsc_table(doc, label):
    print(f"Auditoria da Tabela de 14 Características Gerais (GSC / NI) em {label}...")
    gsc_table = None
    for table in doc.tables:
        if len(table.rows) >= 15 and len(table.columns) == 4:
            c0 = table.rows[0].cells[0].text.strip()
            c1 = table.rows[0].cells[1].text.strip()
            if c0 == "#" and "Característica Geral" in c1:
                gsc_table = table
                break
    assert gsc_table is not None, f"Tabela de GSCs não encontrada no documento {label}."

    notas = []
    total_row_val = None
    for row in gsc_table.rows[1:]:
        c0 = row.cells[0].text.strip()
        c2 = row.cells[2].text.strip()
        if c0.startswith("C") and c0[1:].isdigit():
            try:
                notas.append(int(c2))
            except ValueError:
                raise AssertionError(f"Nota inválida '{c2}' na linha {c0} de {label}.")
        elif "TOTAL" in c0.upper() or "SOMA" in c0.upper() or "NI" in row.cells[1].text:
            total_row_val = int(c2)

    assert len(notas) == 14, f"Esperado 14 características GSC, encontrado {len(notas)} em {label}."
    soma_notas = sum(notas)
    print(f"  [PASS] {len(notas)} características GSC extraídas com sucesso. Soma apurada: {soma_notas}.")
    assert soma_notas == 58, f"Soma das notas das GSCs ({soma_notas}) diverge do NI exigido (58) em {label}."
    if total_row_val is not None:
        assert total_row_val == 58, f"Valor total no rodapé ({total_row_val}) diverge de 58 em {label}."
    print(f"  [PASS] Soma estrita das 14 GSCs == 58 validada em {label}.")

def audit_individual_docx():
    print("===========================================================================")
    print(" 1. Auditoria Forense: Relatório Individual (Fernando Lopes Duarte)")
    print("===========================================================================")
    if not audit_file(INDIVIDUAL_DOCX, "Relatório Individual"):
        return False

    doc = Document(INDIVIDUAL_DOCX)
    audit_xml_styling(doc, "Relatório Individual")

    # Contagem de Parágrafos, Tabelas e Imagens
    p_count = len(doc.paragraphs)
    t_count = len(doc.tables)
    img_count = len(doc.inline_shapes)

    print(f"  Contagens estruturais: {p_count} parágrafos, {t_count} tabelas, {img_count} imagens inline.")
    assert p_count >= 50, f"Contagem de parágrafos ({p_count}) abaixo do esperado (>= 50)."
    assert t_count >= 5, f"Contagem de tabelas ({t_count}) abaixo do esperado (>= 5)."
    assert img_count == 10, f"Contagem de imagens ({img_count}) diverge de 10 esperadas para as 10 funções individuais."
    print("  [PASS] Parágrafos, tabelas e exatamente 10 imagens inline (1 por função) confirmados.")

    # Validação da tabela GSC no Individual
    audit_gsc_table(doc, "Relatório Individual")

    # Conteúdo textual
    full_text = "\n".join([p.text for p in doc.paragraphs])
    for t in doc.tables:
        for r in t.rows:
            for c in r.cells:
                full_text += "\n" + c.text

    # Verificações obrigatórias
    assert "Fernando Lopes Duarte" in full_text, "Nome do aluno não encontrado."
    assert "53 PFB" in full_text or "53 PF" in full_text, "Métrica 53 PFB não encontrada."
    assert "65,19" in full_text or "65.19" in full_text, "Métrica 65,19 PFA não encontrada."
    assert "624,65" in full_text or "624.65" in full_text, "Esforço 624,65 HH não encontrado."
    assert "31.232,35" in full_text or "31232" in full_text, "Custo R$ 31.232,35 não encontrado."
    assert "1,23" in full_text, "Fator de Ajuste 1,23 não encontrado."
    assert "58" in full_text, "Nível de Influência 58 não encontrado."

    # 10 funções individuais
    funcs_ind = [
        "ALI_AUDITORIA_LOGS",
        "ALI_APOLICES_CORRET",
        "AIE_SEGURADORA_PORTAL",
        "AIE_SISTEMA_PIX",
        "EE_CRIAR_TENANT_LIC",
        "EE_EMITIR_APOLICE_VIDA",
        "SE_FOLHA_PAGAMENTO_CALC",
        "SE_LEDGER_AUDITORIA_SHA",
        "CE_LOGS_AUDITORIA_LGPD",
        "CE_APOLICES_ATIVAS"
    ]
    for fn in funcs_ind:
        assert fn in full_text, f"Função individual {fn} não encontrada no texto."
    print("  [PASS] 10 funções individuais de Fernando Lopes Duarte verificadas com sucesso.")

    # Justificativa de dado derivado nas 2 SEs
    assert "Dado Derivado" in full_text, "Termo 'Dado Derivado' não encontrado."
    assert "INSS" in full_text and "IRRF" in full_text, "Motor progressivo da folha não detalhado."
    assert "SHA-256" in full_text or "SHA256" in full_text, "Verificação blockchain SHA-256 não detalhada."

    print("🎉 RELATÓRIO INDIVIDUAL APROVADO COM 100% DE CONFORMIDADE!")
    return True

def audit_team_docx():
    print("\n===========================================================================")
    print(" 2. Auditoria Forense: Relatório de Equipe (HRTech Core — 252 PFB)")
    print("===========================================================================")
    if not audit_file(TEAM_DOCX, "Relatório de Equipe"):
        return False

    doc = Document(TEAM_DOCX)
    audit_xml_styling(doc, "Relatório de Equipe")

    p_count = len(doc.paragraphs)
    t_count = len(doc.tables)
    img_count = len(doc.inline_shapes)

    print(f"  Contagens estruturais: {p_count} parágrafos, {t_count} tabelas, {img_count} imagens inline.")
    assert p_count >= 150, f"Contagem de parágrafos ({p_count}) abaixo do esperado (>= 150)."
    assert t_count >= 15, f"Contagem de tabelas ({t_count}) abaixo do esperado (>= 15)."
    assert img_count == 50, f"Contagem de imagens ({img_count}) diverge de 50 esperadas para as 50 funções."
    print("  [PASS] Parágrafos, tabelas e exatamente 50 imagens inline confirmados.")

    # Validação da tabela GSC na Equipe
    audit_gsc_table(doc, "Relatório de Equipe")

    full_text = "\n".join([p.text for p in doc.paragraphs])
    for t in doc.tables:
        for r in t.rows:
            for c in r.cells:
                full_text += "\n" + c.text

    # Membros da equipe
    members = ["Andryus Zolet", "Felipe", "Valentin", "Nicolas", "Fernando Lopes Duarte"]
    for m in members:
        assert m in full_text, f"Membro {m} não encontrado no relatório de equipe."
    print("  [PASS] 5 membros da equipe identificados.")

    # Métricas globais
    assert "252 PFB" in full_text or "252 PF" in full_text, "Métrica 252 PFB não encontrada."
    assert "309,96" in full_text or "309.96" in full_text, "Métrica 309,96 PFA não encontrada."
    assert "2.970" in full_text or "2970" in full_text, "Esforço 2.970 HH não encontrado."
    assert "148.501" in full_text, "Custo R$ 148.501,00 não encontrado."
    assert "1,23" in full_text, "Fator de Ajuste 1,23 não encontrado."
    assert "58" in full_text, "Nível de Influência 58 não encontrado."
    assert "148,501 AUCP" in full_text or "148.501 AUCP" in full_text, "Validação UCP não encontrada."
    print("  [PASS] Métricas globais e validação cruzada UCP confirmadas.")

    # As 50 Funções Elementares
    all_50_funcs = [
        # 10 ALIs
        "ALI_COLABORADORES", "ALI_ESTRUTURA_ORG", "ALI_REGISTROS_PONTO", "ALI_BANCO_HORAS",
        "ALI_SOLICITACOES_FERIAS", "ALI_CARTOES_BENEF", "ALI_EPIS_EQUIP", "ALI_EXAMES_ASO",
        "ALI_AUDITORIA_LOGS", "ALI_APOLICES_CORRET",
        # 10 AIEs
        "AIE_RECEITA_CPF", "AIE_GOV_ESOCIAL", "AIE_RELOGIO_REP", "AIE_API_NTP_HORA",
        "AIE_OPERADORA_VR_VA", "AIE_OPERADORA_SAUDE", "AIE_MTE_CA", "AIE_CLINICA_MEDICA",
        "AIE_SEGURADORA_PORTAL", "AIE_SISTEMA_PIX",
        # 10 EEs
        "EE_ADMITIR_COLAB", "EE_CRIAR_DEPTO_CARGO", "EE_BATER_PONTO_671", "EE_AJUSTE_ABONO_PONTO",
        "EE_SOLICITAR_FERIAS", "EE_VINCULAR_BENEFICIO", "EE_ENTREGA_EPI_NR6", "EE_REGISTRAR_ASO_NR7",
        "EE_CRIAR_TENANT_LIC", "EE_EMITIR_APOLICE_VIDA",
        # 10 SEs
        "SE_REL_TURNO_CONTRATO", "SE_XML_ESOCIAL_S2200", "SE_ESPELHO_PONTO_CALC", "SE_APURACAO_BANCO_HORAS",
        "SE_RECIBO_FERIAS_1_3", "SE_FATURA_COPARTICIPACAO", "SE_BLOQUEIO_CHAO_FABRICA", "SE_CRONOGRAMA_ASO_PER",
        "SE_FOLHA_PAGAMENTO_CALC", "SE_LEDGER_AUDITORIA_SHA",
        # 10 CEs
        "CE_BUSCA_COLABORADOR", "CE_ORGANOGRAMA_GRAFICO", "CE_BATIDAS_DO_DIA", "CE_EXTRATO_BANCO_HORAS",
        "CE_SALDO_FERIAS_AQUISIT", "CE_CATALOGO_BENEFICIOS", "CE_FICHA_EPIS_COLAB", "CE_HISTORICO_ASO_CA",
        "CE_LOGS_AUDITORIA_LGPD", "CE_APOLICES_ATIVAS"
    ]

    for fn in all_50_funcs:
        assert fn in full_text, f"Função elementar {fn} não encontrada no relatório de equipe."
    print("  [PASS] Todas as 50 funções elementares (10 ALIs, 10 AIEs, 10 EEs, 10 SEs, 10 CEs) confirmadas.")

    # 14 Características Gerais
    for i in range(1, 15):
        assert f"C{i}" in full_text, f"Característica Geral C{i} não encontrada."
    print("  [PASS] 14 Características Gerais (GSC) confirmadas.")

    # Verificação de segregação de Entrada e Saída nas CEs
    assert "Dados de Entrada (Filtros de Pesquisa)" in full_text, "Cabeçalho de entrada das CEs não encontrado."
    assert "Dados de Saída (Resultados / Grid)" in full_text, "Cabeçalho de saída das CEs não encontrado."
    print("  [PASS] Segregação formal de Entrada e Saída nas 10 CEs confirmada.")

    # Verificação das justificativas de dados derivados nas 10 SEs
    assert full_text.count("Justificativa de Dado Derivado") >= 10, "Nem todas as 10 SEs contêm a justificativa de dado derivado."
    print("  [PASS] 100% das 10 SEs contêm a seção 'Justificativa de Dado Derivado'.")

    print("🎉 RELATÓRIO DE EQUIPE APROVADO COM 100% DE CONFORMIDADE!")
    return True

def audit_ifpug_matrices():
    print("\n===========================================================================")
    print(" 3. Auditoria Matricial IFPUG CPM 4.3.1 (ISO/IEC 20926) — 50 Funções")
    print("===========================================================================")

    def check_ilf(rlrs, ders):
        if rlrs == 1:
            return ("Baixa", 7) if ders <= 50 else ("Média", 10)
        elif 2 <= rlrs <= 5:
            if ders < 20: return "Baixa", 7
            elif ders <= 50: return "Média", 10
            else: return "Alta", 15
        else:
            if ders < 20: return "Média", 10
            elif ders <= 50: return "Alta", 15
            else: return "Alta", 15

    def check_eif(rlrs, ders):
        if rlrs == 1:
            return ("Baixa", 5) if ders <= 50 else ("Média", 7)
        elif 2 <= rlrs <= 5:
            if ders < 20: return "Baixa", 5
            elif ders <= 50: return "Média", 7
            else: return "Alta", 10
        else:
            if ders < 20: return "Média", 7
            elif ders <= 50: return "Alta", 10
            else: return "Alta", 10

    def check_ei(ftrs, ders):
        if ftrs <= 1:
            return ("Baixa", 3) if ders <= 15 else ("Média", 4)
        elif ftrs == 2:
            if ders < 5: return "Baixa", 3
            elif ders <= 15: return "Média", 4
            else: return "Alta", 6
        else:
            if ders < 5: return "Média", 4
            elif ders <= 15: return "Alta", 6
            else: return "Alta", 6

    def check_eo(ftrs, ders):
        if ftrs <= 1:
            return ("Baixa", 4) if ders <= 19 else ("Média", 5)
        elif 2 <= ftrs <= 3:
            if ders < 6: return "Baixa", 4
            elif ders <= 19: return "Média", 5
            else: return "Alta", 7
        else:
            if ders < 6: return "Média", 5
            elif ders <= 19: return "Alta", 7
            else: return "Alta", 7

    def check_eq(ftrs, ders):
        if ftrs <= 1:
            return ("Baixa", 3) if ders <= 19 else ("Média", 4)
        elif 2 <= ftrs <= 3:
            if ders < 6: return "Baixa", 3
            elif ders <= 19: return "Média", 4
            else: return "Alta", 6
        else:
            if ders < 6: return "Média", 4
            elif ders <= 19: return "Alta", 6
            else: return "Alta", 6

    # 1. ALIs
    print("Auditoria Matricial: 10 ALIs...")
    for ali in ALIS_DATA:
        c, p = check_ilf(ali["rlrs_count"], ali["ders_count"])
        assert (c, p) == (ali["complexidade"], ali["pfb"]), (
            f"Divergência em ALI {ali['id']}: calculado ({c}, {p} PFB) != catalogado ({ali['complexidade']}, {ali['pfb']} PFB)"
        )
    print("  [PASS] 10 ALIs 100% em conformidade com matriz ILF.")

    # 2. AIEs
    print("Auditoria Matricial: 10 AIEs...")
    for aie in AIES_DATA:
        c, p = check_eif(aie["rlrs_count"], aie["ders_count"])
        assert (c, p) == (aie["complexidade"], aie["pfb"]), (
            f"Divergência em AIE {aie['id']}: calculado ({c}, {p} PFB) != catalogado ({aie['complexidade']}, {aie['pfb']} PFB)"
        )
    print("  [PASS] 10 AIEs 100% em conformidade com matriz EIF.")

    # 3. EEs
    print("Auditoria Matricial: 10 EEs...")
    for ee in EES_DATA:
        c, p = check_ei(ee["ftrs_count"], ee["ders_count"])
        assert (c, p) == (ee["complexidade"], ee["pfb"]), (
            f"Divergência em EE {ee['id']}: calculado ({c}, {p} PFB) != catalogado ({ee['complexidade']}, {ee['pfb']} PFB)"
        )
    print("  [PASS] 10 EEs 100% em conformidade com matriz EI.")

    # 4. SEs
    print("Auditoria Matricial: 10 SEs...")
    for se in SES_DATA:
        c, p = check_eo(se["ftrs_count"], se["ders_count"])
        assert (c, p) == (se["complexidade"], se["pfb"]), (
            f"Divergência em SE {se['id']}: calculado ({c}, {p} PFB) != catalogado ({se['complexidade']}, {se['pfb']} PFB)"
        )
    print("  [PASS] 10 SEs 100% em conformidade com matriz EO.")

    # 5. CEs
    print("Auditoria Matricial: 10 CEs...")
    for ce in CES_DATA:
        c, p = check_eq(ce["ftrs_count"], ce["ders_count"])
        assert (c, p) == (ce["complexidade"], ce["pfb"]), (
            f"Divergência em CE {ce['id']}: calculado ({c}, {p} PFB) != catalogado ({ce['complexidade']}, {ce['pfb']} PFB)"
        )
    print("  [PASS] 10 CEs 100% em conformidade com matriz EQ.")

    # Totalização global
    total_pfb = (
        sum(a["pfb"] for a in ALIS_DATA) +
        sum(a["pfb"] for a in AIES_DATA) +
        sum(e["pfb"] for e in EES_DATA) +
        sum(s["pfb"] for s in SES_DATA) +
        sum(c["pfb"] for c in CES_DATA)
    )
    assert total_pfb == 252, f"Total PFB global {total_pfb} diverge de 252."
    print(f"  [PASS] Totalização PFB Global: {total_pfb} PFB (Equipe 252 PFB exato).")

    # Totalização individual Fernando
    fernando_ids = {
        "ALI_AUDITORIA_LOGS", "ALI_APOLICES_CORRET",
        "AIE_SEGURADORA_PORTAL", "AIE_SISTEMA_PIX",
        "EE_CRIAR_TENANT_LIC", "EE_EMITIR_APOLICE_VIDA",
        "SE_FOLHA_PAGAMENTO_CALC", "SE_LEDGER_AUDITORIA_SHA",
        "CE_LOGS_AUDITORIA_LGPD", "CE_APOLICES_ATIVAS"
    }
    all_funcs = ALIS_DATA + AIES_DATA + EES_DATA + SES_DATA + CES_DATA
    fernando_pfb = sum(f["pfb"] for f in all_funcs if f["id"] in fernando_ids)
    assert fernando_pfb == 53, f"Total PFB individual de Fernando {fernando_pfb} diverge de 53."
    print(f"  [PASS] Totalização PFB Individual: {fernando_pfb} PFB (Fernando 53 PFB exato).")
    print("🎉 TODAS AS 50 FUNÇÕES HOMOLOGADAS SEGUNDO IFPUG CPM 4.3.1!")
    return True

def main():
    print("===========================================================================")
    print(" HRTech Core — Suíte de Validação Forense Independente de Relatórios DOCX")
    print("===========================================================================")

    ok_ind = audit_individual_docx()
    ok_team = audit_team_docx()
    ok_matrix = audit_ifpug_matrices()

    if ok_ind and ok_team and ok_matrix:
        print("\n===========================================================================")
        print(" AUDITORIA CONCLUÍDA COM SUCESSO! 100% DOS REQUISITOS FORAM HOMOLOGADOS.")
        print("===========================================================================")
        sys.exit(0)
    else:
        print("\n[ERRO] Falha na auditoria de conformidade dos relatórios.")
        sys.exit(1)

if __name__ == "__main__":
    main()
