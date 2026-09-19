#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Script de Geração do Relatório Acadêmico em Word (.docx)
Disciplina: Medição e Análise de Software — PUCPR
Projeto: HRTech Core — Padrões de Projeto, Banco de Dados SQLite e Linha de Produção de Software (LPS)
Autores (5 Integrantes): Fernando Lopes Duarte, Andryus, Felipe, Valentin, Nicholas
"""

import os
import sys
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def set_table_borders(table, color="D1D5DB", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideV w:val="none"/>'
        f'<w:left w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def format_cell(cell, text, bold=False, italic=False, color=RGBColor(31, 41, 55), font_size=9.5, align=WD_ALIGN_PARAGRAPH.LEFT, bg_color=None):
    if bg_color:
        set_cell_background(cell, bg_color)
    set_cell_margins(cell, top=90, bottom=90, left=120, right=120)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = Pt(12)
    run = p.add_run(str(text))
    run.font.name = 'Arial'
    run.font.size = Pt(font_size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color

def add_heading_1(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(16)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = 'Arial'
    run.font.size = Pt(14)
    run.font.bold = True
    run.font.color.rgb = RGBColor(30, 58, 138) # Deep Blue #1E3A8A

def add_heading_2(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = 'Arial'
    run.font.size = Pt(11.5)
    run.font.bold = True
    run.font.color.rgb = RGBColor(17, 24, 39) # Dark Slate #111827

def add_heading_3(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = 'Arial'
    run.font.size = Pt(10.5)
    run.font.bold = True
    run.font.color.rgb = RGBColor(55, 65, 81) # Gray #374151

def add_body_p(doc, text, bold_prefix="", italic=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_bold = p.add_run(bold_prefix)
        r_bold.font.name = 'Arial'
        r_bold.font.size = Pt(10)
        r_bold.font.bold = True
        r_bold.font.color.rgb = RGBColor(31, 41, 55)
    run = p.add_run(text)
    run.font.name = 'Arial'
    run.font.size = Pt(10)
    run.font.italic = italic
    run.font.color.rgb = RGBColor(55, 65, 81)
    return p

def add_callout_box(doc, text, title="NOTA DA EQUIPE"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tblPr = tbl._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="none"/>'
        f'<w:bottom w:val="none"/>'
        f'<w:left w:val="single" w:sz="24" w:space="0" w:color="2563EB"/>'
        f'<w:right w:val="none"/>'
        f'<w:insideH w:val="none"/>'
        f'<w:insideV w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)
    cell = tbl.rows[0].cells[0]
    set_cell_background(cell, "EFF6FF")
    set_cell_margins(cell, top=100, bottom=100, left=150, right=150)
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    
    r_title = p.add_run(f"📌 {title}: ")
    r_title.font.name = 'Arial'
    r_title.font.size = Pt(9.5)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(30, 64, 175)
    
    r_txt = p.add_run(text)
    r_txt.font.name = 'Arial'
    r_txt.font.size = Pt(9.5)
    r_txt.font.color.rgb = RGBColor(30, 58, 138)
    
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def add_code_block(doc, code_str, language="php"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl, color="E2E8F0", sz="4")
    cell = tbl.rows[0].cells[0]
    set_cell_background(cell, "F8FAFC")
    set_cell_margins(cell, top=90, bottom=90, left=140, right=140)
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = Pt(11)
    run = p.add_run(code_str)
    run.font.name = 'Courier New'
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor(15, 23, 42)
    
    p_sp = doc.add_paragraph()
    p_sp.paragraph_format.space_before = Pt(0)
    p_sp.paragraph_format.space_after = Pt(4)

def build_student_delivery_docx(output_path):
    doc = Document()

    # Configuração de Margens (2,5 cm = ~0.98 in)
    for section in doc.sections:
        section.top_margin = Inches(0.98)
        section.bottom_margin = Inches(0.98)
        section.left_margin = Inches(0.98)
        section.right_margin = Inches(0.98)

    # ---------------------------------------------------------
    # CABEÇALHO ACADÊMICO ESTILO ALUNO
    # ---------------------------------------------------------
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(2)
    run_t = p_title.add_run("Relatório de Entrega: Padrões de Projeto, Persistência Relacional e Linha de Produção de Software (LPS)")
    run_t.font.name = 'Arial'
    run_t.font.size = Pt(16)
    run_t.font.bold = True
    run_t.font.color.rgb = RGBColor(30, 58, 138)

    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_after = Pt(8)
    run_s = p_sub.add_run("Projeto: HRTech Core — Plataforma Modular de Recursos Humanos (PUCPR)")
    run_s.font.name = 'Arial'
    run_s.font.size = Pt(11)
    run_s.font.bold = True
    run_s.font.color.rgb = RGBColor(75, 85, 99)

    # Info Box Acadêmica
    tbl_info = doc.add_table(rows=4, cols=2)
    tbl_info.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl_info, color="E5E7EB", sz="4")

    info_meta = [
        ("Instituição / Escola:", "Pontifícia Universidade Católica do Paraná (PUCPR) — Escola Politécnica"),
        ("Disciplina / Semestre:", "Medição e Análise de Software (2026-2)"),
        ("Equipe de Alunos (5 membros):", "Fernando Lopes Duarte, Andryus, Felipe, Valentin, Nicholas"),
        ("Escopo da Entrega:", "Item 01 (12 Entidades + 8 Padrões) e Item 02 (Banco SQLite 12 tabelas, 10 CRUDs e LPS)")
    ]

    for r_idx, (lab, val) in enumerate(info_meta):
        row = tbl_info.rows[r_idx]
        format_cell(row.cells[0], lab, bold=True, color=RGBColor(30, 58, 138), bg_color="F3F4F6", font_size=9)
        format_cell(row.cells[1], val, bold=False, color=RGBColor(31, 41, 55), bg_color="F9FAFB", font_size=9)
        row.cells[0].width = Inches(2.3)
        row.cells[1].width = Inches(4.7)

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # Identificação da Equipe e Atribuição de Módulos
    add_heading_2(doc, "Equipe de Engenharia (5 Integrantes) e Divisão de Responsabilidades")
    
    tbl_team = doc.add_table(rows=6, cols=3)
    tbl_team.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl_team, color="D1D5DB", sz="4")

    team_headers = ["Integrante", "Papel Principal no Projeto", "Módulos & CRUDs Sob Responsabilidade"]
    for c_idx, th in enumerate(team_headers):
        format_cell(tbl_team.rows[0].cells[c_idx], th, bold=True, color=RGBColor(255, 255, 255), bg_color="1E3A8A", font_size=9)

    team_data = [
        ("Fernando Lopes Duarte", "Arquiteto Core, Database Lead & LPS Engineer", "CRUD 1: Gestão de Tenants (Empresas)\nCRUD 2: Gestão de Usuários e RBAC"),
        ("Andryus", "Engenheiro de Domínio & Estrutura Organizacional", "CRUD 3: Cadastro de Colaboradores\nCRUD 4: Gestão de Cargos e Departamentos"),
        ("Felipe", "Engenheiro de Compliance & Ponto Eletrônico", "CRUD 5: Registro de Ponto (Portaria 671 & SHA-256)\nCRUD 6: Gestão de Ajustes e Abonos de Ponto"),
        ("Valentin", "Engenheiro de Benefícios & Direitos Trabalhistas", "CRUD 7: Gestão e Agendamento de Férias CLT\nCRUD 8: Catálogo e Manutenção de Benefícios"),
        ("Nicholas", "Engenheiro de SST & Integração de Seguros", "CRUD 9: Controle de EPIs (NR-6) e Exames ASO (NR-7)\nCRUD 10: Apólices do Portal do Corretor FinCorp")
    ]

    for r_idx, row_values in enumerate(team_data, start=1):
        bg = "F9FAFB" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(row_values):
            format_cell(tbl_team.rows[r_idx].cells[c_idx], val, bold=(c_idx==0), bg_color=bg, font_size=8.5)

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # ---------------------------------------------------------
    # SEÇÃO 1: APRESENTAÇÃO E CONTEXTO
    # ---------------------------------------------------------
    add_heading_1(doc, "1. Apresentação e Contexto da Entrega")
    add_body_p(doc, "Professor, a nossa equipe de 5 alunos desenvolveu a arquitetura de backend e o relatório técnico da plataforma HRTech Core, cumprindo integralmente os requisitos dos slides da disciplina:")
    add_body_p(doc, "1. Item 01 do Slide — Padrões de Projeto (Opção 01): Modelagem rigorosa de 12 Entidades de Domínio tipadas em PHP 8.3.6, apoiadas por Value Objects imutáveis (CNPJ, CPF com máscara LGPD, Centavos Inteiros com Fowler Allocation e Geolocalização Haversine) e 8 Padrões de Projeto operacionais (2 Singletons, 3 Template Methods e 3 Strategies).")
    add_body_p(doc, "2. Item 02 do Slide — Banco de Dados Relacional, 10 CRUDs e Variabilidade LPS: Camada SQLite com 12 tabelas, 11 índices de performance, chaves estrangeiras com cascata e transações ACID; 10 operações CRUD distribuídas entre os 5 alunos (2 por integrante); e o motor de Linha de Produção de Software (LPS) demonstrando regras de negócio altamente variáveis entre os perfis de Tecnologia, Indústria e Financeiro.")
    add_body_p(doc, "A plataforma conta com uma suíte de 868 asserções de testes unitários e de integração com 100% de aprovação, testes adversariais de estresse contra injeção SQL e quebra de regras da LPS, um runner de linha de comando (run_demo.php) e um protótipo web com 17 telas.")

    # ---------------------------------------------------------
    # SEÇÃO 2: ITEM 01 DO SLIDE
    # ---------------------------------------------------------
    add_heading_1(doc, "2. Item 01 do Slide: 12 Classes de Domínio & 8 Padrões de Projeto (Opção 01)")
    
    add_heading_2(doc, "2.1 As 12 Entidades de Domínio e Value Objects")
    add_body_p(doc, "As entidades de domínio foram estruturadas em src/Domain/Entities/ aplicando os princípios de Clean Architecture e DDD. A tabela a seguir especifica cada entidade e suas responsabilidades:")

    tbl_ent = doc.add_table(rows=13, cols=3)
    tbl_ent.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl_ent, color="D1D5DB", sz="4")

    ent_headers = ["#", "Entidade de Domínio", "Descrição Funcional e Regras de Negócio"]
    for c_idx, eh in enumerate(ent_headers):
        format_cell(tbl_ent.rows[0].cells[c_idx], eh, bold=True, color=RGBColor(255, 255, 255), bg_color="1E3A8A", font_size=9)

    entities_data = [
        ("01", "Tenant", "Empresa cliente na modalidade multi-tenant. Contém CNPJ validado, razão social, nome fantasia, segmento de mercado (Tech, Indústria, Financeiro) e licenças ativas."),
        ("02", "User", "Credencial de acesso e autenticação com hash bcrypt seguro, RBAC com matriz de permissões, suporte a MFA e associação com colaborador."),
        ("03", "Employee", "Colaborador corporativo com CPF validado, regime contratual (CLT, PJ, Estagiário), remuneração em centavos inteiros, saldo de férias e saldo de banco de horas em minutos."),
        ("04", "Department", "Unidade organizacional interna com código alfanumérico único por tenant, nome, centro de custo e vinculação do gestor responsável."),
        ("05", "Role", "Cargo corporativo com nível hierárquico numérico de 1 a 100 (utilizado para aprovação de alçadas) e permissões de acesso em JSON."),
        ("06", "TimeLog", "Registro eletrônico de ponto sob a Portaria 671/2021 MTE. Possui NSR sequencial, carimbo de tempo ISO 8601, geolocalização e hash SHA-256 encadeado. Imutável por design."),
        ("07", "TimeAdjustmentRequest", "Solicitação de retificação ou inclusão manual de ponto com justificativa, comprovante e fluxo formal de aprovação pelo gestor (PENDING, APPROVED, REJECTED)."),
        ("08", "VacationRequest", "Solicitação de férias sob os Artigos 129 a 145 da CLT. Valida período mínimo de 5 dias corridos, abono pecuniário (venda de 1/3) e adiantamento do 13º salário."),
        ("09", "Benefit", "Pacote de benefícios (Vale-Transporte, Vale-Alimentação, Plano de Saúde) com valor mensal em centavos e percentual de coparticipação do colaborador."),
        ("10", "EquipmentASO", "Gestão de EPIs entregues com Certificado de Aprovação (NR-6) e controle de exames médicos ocupacionais ASO (NR-7) com parecer Apto/Inapto e CRM do médico."),
        ("11", "InsurancePolicy", "Apólice de seguro de vida corporativo e seguro D&O operados no Portal do Corretor FinCorp Seguros, com prêmio mensal e capital total segurado."),
        ("12", "AuditLog", "Trilha de auditoria criptográfica imutável para atendimento à LGPD e SOX. Registra ator, ação, deltas de estado e assinatura SHA-256 encadeada.")
    ]

    for r_idx, (num, name, desc) in enumerate(entities_data, start=1):
        bg = "F9FAFB" if r_idx % 2 == 1 else "FFFFFF"
        format_cell(tbl_ent.rows[r_idx].cells[0], num, bold=True, bg_color=bg, font_size=8.5, align=WD_ALIGN_PARAGRAPH.CENTER)
        format_cell(tbl_ent.rows[r_idx].cells[1], name, bold=True, color=RGBColor(30, 58, 138), bg_color=bg, font_size=8.5)
        format_cell(tbl_ent.rows[r_idx].cells[2], desc, bg_color=bg, font_size=8.5)

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    add_callout_box(doc, "Para blindar a integridade dos dados, implementamos 4 Value Objects imutáveis: Cnpj e Cpf (validação módulo 11 e máscara LGPD), Money (centavos inteiros imunes a erros de arredondamento IEEE 754 de ponto flutuante com partição justa de Martin Fowler) e GeoLocation (cálculo de cerca eletrônica via fórmula de Haversine).", "INTEGRIDADE POR VALUE OBJECTS")

    # Os 8 Padrões de Projeto
    add_heading_2(doc, "2.2 Os 8 Padrões de Projeto Implementados (Opção 01)")
    add_body_p(doc, "A equipe consolidou a implementação dos 8 padrões da Opção 01 solicitada no slide:")

    add_heading_3(doc, "A) 2 Padrões Singleton")
    add_body_p(doc, "1. TenantContextManager (src/Patterns/Singleton/TenantContextManager.php): Garante que a requisição ativa opere estritamente no escopo da empresa cliente correta, impedindo vazamento cruzado de dados (cross-tenant leakage). Possui o método runInContext() que gerencia escopos temporários com restauração garantida em bloco finally.")
    add_body_p(doc, "2. AuditLogger (src/Patterns/Singleton/AuditLogger.php): Cria um livro-razão imutável de eventos corporativos encadeados criptograficamente via SHA-256. Cada registro inclui o hash do evento anterior (previous_hash), permitindo validação integral da cadeia através do método verifyChainIntegrity().")

    add_heading_3(doc, "B) 3 Padrões Template Method")
    add_body_p(doc, "1. PayrollCalculatorTemplate (src/Patterns/TemplateMethod/Payroll/): Define o algoritmo macro de fechamento da folha através do método final calculatePayroll(). As etapas invariantes (validação de elegibilidade e montagem do holerite) são fixas, enquanto os passos variantes são delegados às subclasses concretas: CltPayroll (tabelas progressivas oficiais de INSS e IRRF), PjPayroll (retenções fiscais empresariais IRRF 1,5% e CSRF 4,65%) e InternPayroll (isenção de encargos trabalhistas sob a Lei 11.788/2008).")
    add_body_p(doc, "2. TimeLogImporterTemplate (src/Patterns/TemplateMethod/Importer/): Padroniza o fluxo de ingestão de marcações de ponto com os métodos CsvImporter (arquivos legados com mapeamento de sinônimos), JsonImporter (webhooks móveis) e ApiImporter (relógios de ponto homologados REP com verificação de Bearer Token).")
    add_body_p(doc, "3. ReportGeneratorTemplate (src/Patterns/TemplateMethod/Report/): Orquestra a extração, isolamento por tenant, cabeçalho, paginação e exportação nos formatos PdfReportGenerator (relatório textual paginado), ExcelReportGenerator (leiaute tabular com totalizadores de soma exata) e JsonReportGenerator (metadados analíticos estruturados).")

    add_heading_3(doc, "C) 3 Padrões Strategy")
    add_body_p(doc, "1. OvertimeStrategy (src/Patterns/Strategy/Overtime/): Implementa o cálculo legal de horas extras com Standard50Strategy (acréscimo de 50% em dias úteis conforme Art. 59 §1 da CLT), Sunday100Strategy (acréscimo de 100% em domingos e feriados conforme Súmula 146 do TST) e BankHoursStrategy (remuneração financeira R$ 0,00 e crédito de minutos compensatórios no saldo do colaborador conforme Art. 59 §2 da CLT).")
    add_body_p(doc, "2. BenefitDiscountStrategy (src/Patterns/Strategy/BenefitDiscount/): Descontos salariais de coparticipação com TransportationVoucherStrategy (limite legal estrito de até 6% do salário-base conforme Lei 7.418/1985), HealthPlanStrategy (coparticipação escalonada por faixas etárias da ANS) e MealVoucherStrategy (coparticipação limitada a 20% pelo Programa de Alimentação do Trabalhador - PAT).")
    add_body_p(doc, "3. PerformanceBonusStrategy (src/Patterns/Strategy/Performance/): Bonificação flexível com OkrStrategy (metas ágeis trimestrais com overachievement de até 120%), Evaluation360Strategy (média balanceada entre gestor, autoavaliação e pares para ambientes industriais) e KpiStrategy (indicadores numéricos com cláusula de corte mínima, métricas inversas e bônus agressivo para o setor financeiro).")

    # ---------------------------------------------------------
    # SEÇÃO 3: ITEM 02 DO SLIDE
    # ---------------------------------------------------------
    add_heading_1(doc, "3. Item 02 do Slide: Banco de Dados Relacional, 10 CRUDs e Variabilidade LPS")

    add_heading_2(doc, "3.1 Banco de Dados Relacional SQLite (12 Tabelas e 11 Índices)")
    add_body_p(doc, "A camada de banco de dados foi construída com a classe DatabaseManager (PDO SQLite), operando no arquivo local database.sqlite. O banco é configurado com modo WAL (Write-Ahead Logging) e chaves estrangeiras ativas (PRAGMA foreign_keys = ON;). A integridade transacional é assegurada pelo método transaction(callable), com rollback automático em falhas.")
    add_body_p(doc, "Foram estruturadas 12 tabelas relacionais completas: tenants, departments, roles, employees, users, time_logs, time_adjustment_requests, vacation_requests, benefits, equipment_aso, insurance_policies e audit_logs. Para suportar a arquitetura multi-tenant com alta performance, foram definidos 11 índices estratégicos.")

    add_heading_2(doc, "3.2 Mapeamento dos 10 CRUDs pelos 5 Integrantes")
    add_body_p(doc, "Os 10 módulos CRUD foram equitativamente divididos entre os 5 integrantes do grupo, seguindo o fluxo arquitetural Repositório Concreto -> Serviço de Negócio:")

    tbl_cruds = doc.add_table(rows=11, cols=4)
    tbl_cruds.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl_cruds, color="D1D5DB", sz="4")

    crud_headers = ["Integrante", "CRUD", "Módulo / Entidade", "Regras de Negócio e Validações"]
    for c_idx, ch in enumerate(crud_headers):
        format_cell(tbl_cruds.rows[0].cells[c_idx], ch, bold=True, color=RGBColor(255, 255, 255), bg_color="1E3A8A", font_size=8.5)

    cruds_data = [
        ("Fernando Lopes Duarte", "CRUD 1", "Gestão de Tenants", "Validação formal de CNPJ (módulo 11), unicidade no sistema e controle de licenças ativas."),
        ("Fernando Lopes Duarte", "CRUD 2", "Gestão de Usuários", "Hash bcrypt (password_hash), proteção de complexidade de senha, controle MFA e unicidade de username por tenant."),
        ("Andryus", "CRUD 3", "Cadastro de Colaboradores", "Validação formal de CPF, salários armazenados em centavos, saldo inicial de 30 dias de férias e crédito/débito de banco de horas."),
        ("Andryus", "CRUD 4", "Cargos e Departamentos", "Unicidade de código departamental por tenant, centros de custo e níveis de hierarquia de 1 a 100 para aprovações."),
        ("Felipe", "CRUD 5", "Ponto Eletrônico (Portaria 671)", "Sequenciamento de NSR, geolocalização com tolerância, assinatura SHA-256 e imutabilidade estrita (rejeição de UPDATE e DELETE)."),
        ("Felipe", "CRUD 6", "Ajustes de Ponto", "Solicitação formal de inclusão manual de batidas, parecer com justificativa e workflow de aprovação pelo gestor com carimbo."),
        ("Valentin", "CRUD 7", "Solicitações de Férias", "Cumprimento dos Artigos 129 a 145 da CLT, período mínimo de 5 dias, abono pecuniário (1/3) e débito automático no saldo de férias."),
        ("Valentin", "CRUD 8", "Gestão de Benefícios", "Parametrização de benefícios (VT, VR, VA, Saúde), controle de coparticipação percentual e verificação de tetos legais."),
        ("Nicholas", "CRUD 9", "Controle de EPIs e ASO", "Controle de EPI com Certificado de Aprovação (NR-6) e exames ocupacionais admissionais e periódicos com CRM e parecer Apto/Inapto (NR-7)."),
        ("Nicholas", "CRUD 10", "Apólices FinCorp Seguros", "Emissão, renovação e cancelamento de apólices de vida corporativas e D&O, prêmios mensais e controle de capital segurado.")
    ]

    for r_idx, (member, crud_id, mod, rules) in enumerate(cruds_data, start=1):
        bg = "F9FAFB" if r_idx % 2 == 1 else "FFFFFF"
        format_cell(tbl_cruds.rows[r_idx].cells[0], member, bold=True, bg_color=bg, font_size=8)
        format_cell(tbl_cruds.rows[r_idx].cells[1], crud_id, bold=True, color=RGBColor(30, 58, 138), bg_color=bg, font_size=8, align=WD_ALIGN_PARAGRAPH.CENTER)
        format_cell(tbl_cruds.rows[r_idx].cells[2], mod, bold=False, bg_color=bg, font_size=8)
        format_cell(tbl_cruds.rows[r_idx].cells[3], rules, bg_color=bg, font_size=8)

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # Variabilidade LPS
    add_heading_2(doc, "3.3 Motor de Variabilidade da Linha de Produção de Software (LPS)")
    add_body_p(doc, "O motor de variabilidade (LpsVariabilityEngine) personaliza as regras de negócio e módulos da aplicação conforme o perfil de mercado do cliente:")

    tbl_lps = doc.add_table(rows=12, cols=4)
    tbl_lps.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl_lps, color="D1D5DB", sz="4")

    lps_headers = ["Feature Flag", "Finalidade na LPS", "Tech", "Indústria", "Financeiro"]
    # We will adjust to 5 columns
    pass

    # Re-create table with 5 columns
    tbl_lps._element.getparent().remove(tbl_lps._element)
    tbl_lps = doc.add_table(rows=12, cols=5)
    tbl_lps.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl_lps, color="D1D5DB", sz="4")

    lps_headers = ["Feature Flag", "Finalidade na Regra de Negócio", "Tech", "Indústria", "Financeiro"]
    for c_idx, lh in enumerate(lps_headers):
        format_cell(tbl_lps.rows[0].cells[c_idx], lh, bold=True, color=RGBColor(255, 255, 255), bg_color="1E3A8A", font_size=8)

    lps_data = [
        ("bank_of_hours", "Compensação de jornada suplementar em banco", "ATIVO", "Inativo", "Inativo"),
        ("overtime_payout", "Pagamento de horas extras em dinheiro no holerite", "Inativo", "ATIVO", "ATIVO"),
        ("risk_ppe_required", "Exigência de EPIs (NR-6) e Atestado ASO (NR-7)", "Dispensado", "OBRIGATÓRIO", "Dispensado"),
        ("flexible_benefits", "Cartão multibenefícios com saldo flexível", "ATIVO", "Inativo", "Inativo"),
        ("d_and_o_insurance", "Seguro de responsabilidade para diretores", "ATIVO", "Inativo", "Inativo"),
        ("biometric_punch_mandatory", "Exigência estrita de biometria no ponto eletrônico", "Inativo", "Inativo", "OBRIGATÓRIO"),
        ("chartered_transport", "Fornecimento de transporte fretado próprio", "Inativo", "ATIVO", "Inativo"),
        ("executive_health_plan", "Plano médico executivo hospitalar", "Inativo", "Inativo", "ATIVO"),
        ("aggressive_bonus", "Multiplicador de bonificação executiva", "Inativo", "Inativo", "ATIVO"),
        ("strict_lgpd_audit", "Trilha de auditoria estendida para conformidade SOX", "Inativo", "Inativo", "ATIVO"),
        ("fincorp_life_policy", "Apólice de seguro de vida do Portal do Corretor", "Inativo", "Inativo", "ATIVO")
    ]

    for r_idx, row_values in enumerate(lps_data, start=1):
        bg = "F9FAFB" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(row_values):
            is_active = val in ["ATIVO", "OBRIGATÓRIO"]
            txt_color = RGBColor(30, 58, 138) if is_active else RGBColor(55, 65, 81)
            format_cell(tbl_lps.rows[r_idx].cells[c_idx], val, bold=(c_idx==0 or is_active), color=txt_color, bg_color=bg, font_size=8)

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    add_heading_3(doc, "Bloqueio Preventivo de Segurança Ocupacional (Work Eligibility Blocking)")
    add_body_p(doc, "Um dos destaques da LPS é a proteção no chão de fábrica da Indústria. O método validateWorkEligibility() bloqueia sumariamente o início do turno de um colaborador caso:")
    add_body_p(doc, "• O colaborador não possua registro de Atestado de Saúde Ocupacional (código ASO_MISSING);")
    add_body_p(doc, "• O médico tenha considerado o colaborador clinicamente inapto (código ASO_UNFIT);")
    add_body_p(doc, "• O exame médico periódico de ASO estiver vencido sob a NR-7 (código ASO_EXPIRED);")
    add_body_p(doc, "• O Certificado de Aprovação (CA) do EPI entregue estiver vencido sob a NR-6 (código PPE_CA_EXPIRED).")

    add_code_block(doc, """// Bloqueio de turno fabril no LpsVariabilityEngine.php:
if ($aso === null) {
    return ['allowed' => false, 'reason' => 'Trabalho bloqueado: Sem ASO cadastrado.', 'code' => 'ASO_MISSING'];
}
if (!$aso->isFit) {
    return ['allowed' => false, 'reason' => 'Trabalho bloqueado: Colaborador INAPTO.', 'code' => 'ASO_UNFIT'];
}
if ($aso->isExamExpired()) {
    return ['allowed' => false, 'reason' => 'Trabalho bloqueado: ASO vencido (NR-7).', 'code' => 'ASO_EXPIRED'];
}
if ($aso->isCaExpired()) {
    return ['allowed' => false, 'reason' => 'Trabalho bloqueado: EPI com CA vencido (NR-6).', 'code' => 'PPE_CA_EXPIRED'];
}
return ['allowed' => true, 'reason' => 'Colaborador liberado para o turno.', 'code' => 'COMPLIANT'];""")

    # ---------------------------------------------------------
    # SEÇÃO 4: GUIA PASSO A PASSO DE EXECUÇÃO
    # ---------------------------------------------------------
    add_heading_1(doc, "4. Guia Passo a Passo de Execução para Avaliação")

    add_heading_2(doc, "4.1 Execução do Script CLI de Demonstração Unificada (run_demo.php)")
    add_body_p(doc, "Para executar a demonstração completa da plataforma no terminal:")
    add_code_block(doc, "cd hrtech_backend_patterns\nphp run_demo.php")
    add_body_p(doc, "O script recria a base SQLite, popula dados de teste, roda os 8 Design Patterns, executa os 10 CRUDs dos 5 alunos e demonstra a variabilidade LPS nos perfis Tech, Indústria e Financeiro.")

    add_heading_2(doc, "4.2 Execução das Suítes de Testes Automatizados")
    add_body_p(doc, "A plataforma conta com 868 asserções regressivas e 676 asserções adversariais (100% aprovadas):")
    add_code_block(doc, """# Suíte Regressiva Oficial (868 asserções):
php tests/m1_verify.php   # Value Objects, Enums e Contratos (132 asserções)
php tests/m2_verify.php   # 12 Entidades de Domínio e Multi-Tenant (376 asserções)
php tests/m3_verify.php   # 2 Singletons, 3 Templates e 3 Strategies (200 asserções)
php tests/m4_verify.php   # SQLite 12 tabelas, 10 CRUDs e Motor LPS (160 asserções)

# Suíte de Estresse e Adversarial (676 asserções):
php tests/m4_adversarial_cruds.php  # Teste contra SQLi e Foreign Keys (416 asserções)
php tests/m4_adversarial_lps.php    # Fuzzing de 10.000 iterações na LPS (260 asserções)""")

    add_heading_2(doc, "4.3 Acesso ao Protótipo Web Navegável (Porta 8085)")
    add_body_p(doc, "Para inspecionar as 17 telas funcionais do protótipo visual:")
    add_code_block(doc, "# Na raiz do projeto:\npython3 -m http.server 8085\n# Acesse no navegador: http://localhost:8085")
    add_body_p(doc, "No protótipo web, o usuário pode interagir com o organograma, registrar batidas de ponto com geolocalização e alternar os segmentos da LPS em tempo real.")

    # ---------------------------------------------------------
    # SEÇÃO 5: CONCLUSÃO
    # ---------------------------------------------------------
    add_heading_1(doc, "5. Conclusão e Considerações da Equipe")
    add_body_p(doc, "A equipe concluiu o projeto HRTech Core consolidando com rigor os objetivos da disciplina de Medição e Análise de Software. O projeto comprova a aplicação prática de padrões de projeto consolidados (GoF), tipagem moderna do PHP 8.3, persistência transacional relacional e arquitetura de Linha de Produção de Software com alto grau de parametrização e conformidade legal.")
    add_body_p(doc, "Agradecemos ao professor pela oportunidade prática de aprendizado e estamos à disposição para eventuais esclarecimentos técnicos.")

    # Salvamento
    doc.save(output_path)
    print(f"[OK] Documento Word gerado com sucesso em: {output_path}")

if __name__ == '__main__':
    target = sys.argv[1] if len(sys.argv) > 1 else "Entrega_Projeto_Equipe_Padroes_LPS_HRTech.docx"
    build_student_delivery_docx(target)
